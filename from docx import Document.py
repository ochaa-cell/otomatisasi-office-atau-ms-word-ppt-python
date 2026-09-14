
from docx import Document

def buat_dokumen():
    # Membuat dokumen baru
    doc = Document()

    # Menambahkan Judul
    doc.add_heading('Laporan Tugas Microsoft Word Specialist', 0)

    # Menambahkan Paragraf
    doc.add_paragraph('Ini adalah contoh dokumen yang dibuat secara otomatis menggunakan Python. '
                      'Proyek ini bertujuan untuk mendemonstrasikan kemampuan otomatisasi Microsoft Word.')

    # Menambahkan Sub-judul dan list
    doc.add_heading('Poin-poin Penting:', level=1)
    doc.add_paragraph('Mudah dipelajari', style='List Bullet')
    doc.add_paragraph('Sangat cepat', style='List Bullet')
    doc.add_paragraph('Cocok untuk tugas akhir/proyek GitHub', style='List Bullet')

    # Menyimpan dokumen
    doc.save('Tugas_Word_Specialist.docx')
    print("Dokumen Word berhasil dibuat: Tugas_Word_Specialist.docx")

if __name__ == '__main__':
    buat_dokumen()



from pptx import Presentation

def buat_presentasi():
    # Membuat presentasi baru
    prs = Presentation()

    # Slide 1: Slide Judul
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    title = slide1.shapes.title
    subtitle = slide1.placeholders[1]
    title.text = "Tugas PowerPoint Specialist"
    subtitle.text = "Dibuat otomatis dengan Python"

    # Slide 2: Slide Isi dengan Bullet Points
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide2.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Mengapa Otomatisasi?"
    tf = body_shape.text_frame
    tf.text = "Menghemat waktu pengerjaan berulang"
    
    p = tf.add_paragraph()
    p.text = "Meminimalisir kesalahan manusia (human error)"
    p.level = 1
    
    p2 = tf.add_paragraph()
    p2.text = "Sangat keren untuk portofolio GitHub"
    p2.level = 1

    # Menyimpan presentasi
    prs.save('Tugas_PPT_Specialist.pptx')
    print("Presentasi PowerPoint berhasil dibuat: Tugas_PPT_Specialist.pptx")

if __name__ == '__main__':
    buat_presentasi()